import pandas as pd
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.formatting.rule import CellIsRule
from io import BytesIO

def analise_preco(arquivo):

    arquivo['DESCRICAO'] = arquivo['DESCRICAO'].str.strip()

    arquivo = arquivo[arquivo['STATUS'] == "OK"]

    colunas_inicio = ['ESTADO','COD_FABRICANTE', 'DESCRICAO']

    colunas_final = ['MENOR PRECO','PRECO EMPRESA', 'DIF %', 'PRECO CLUBE', 'DIF % CLUBE']

    concorrentes = arquivo['CONCORRENTE'].unique().tolist()

    comparativo = arquivo.pivot_table(index=['ESTADO','COD_FABRICANTE'], 
                                    columns='CONCORRENTE', 
                                    values=['PRECO CONCORRENTE'],
                                    aggfunc='first')

    comparativo = comparativo.rename_axis(columns=[None, None])

    comparativo = comparativo.droplevel(0, axis=1)

    comparativo.reset_index(drop=False, inplace=True)

    comparativo_final = comparativo.merge(
        arquivo[['DESCRICAO', 'PRECO EMPRESA', 'PRECO CLUBE', 'ESTADO', 'COD_FABRICANTE']],
        left_on=['ESTADO', 'COD_FABRICANTE'],
        right_on=['ESTADO', 'COD_FABRICANTE'],
        how='left')

    comparativo_final = comparativo_final.drop_duplicates(
        subset=['ESTADO', 'COD_FABRICANTE']
    )

    comparativo_final['MENOR PRECO'] = comparativo_final.loc[:,concorrentes].min(axis=1)

    comparativo_final['DIF %'] = comparativo_final['PRECO EMPRESA'] / comparativo_final['MENOR PRECO'] -1

    comparativo_final['DIF % CLUBE'] = comparativo_final['PRECO CLUBE'] / comparativo_final['MENOR PRECO'] -1

    ordena_colunas = colunas_inicio + concorrentes + colunas_final

    comparativo_final = comparativo_final[ordena_colunas]
    
    output = BytesIO()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for estado, dados in comparativo_final.groupby('ESTADO'):
            dados.to_excel(
                writer,
                sheet_name= estado,
                index = False)
            
            ws = writer.book[estado]
            
            # Cabeçalho
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(fill_type='solid', fgColor='92D050')
                cell.alignment = Alignment(horizontal='center', vertical='center')
                ws.row_dimensions[1].height = 27
                
            # Bordas
            borda = Side(style='thin', color='000000')
            for row in ws.iter_rows():
                for cell in row:
                    cell.border = Border(left=borda, right=borda, top=borda, bottom=borda)
            # Formatação Moeda  
            colunas_moeda = [*concorrentes,'MENOR PRECO', 'PRECO EMPRESA', 'PRECO CLUBE']
            
            for coluna in colunas_moeda:
                col = comparativo_final.columns.get_loc(coluna) + 1
                
                for cell in ws.iter_cols(
                    min_col=col,
                    max_col=col,
                    min_row=2
                ):
                    for c in cell:
                        c.number_format = 'R$ #,##0.00'
            
            # Formatação Percentual         
            colunas_percentual = ['DIF %', 'DIF % CLUBE']
            
            for coluna in colunas_percentual:
                col = comparativo_final.columns.get_loc(coluna) + 1
                
                for cell in ws.iter_cols(
                    min_col=col,
                    max_col=col,
                    min_row=2
                ):
                    for c in cell:
                        c.number_format = '0.00%'
                        
            # Formatação Condicional
            verde_fill = PatternFill(fill_type='solid', bgColor='C6EFCE')
            verde_font = Font(color='006100')
            
            vermelho_fill = PatternFill(fill_type='solid', bgColor='FFC7CE')
            vermelho_font = Font(color='9C0006')
            
            for coluna in colunas_percentual:
                col = comparativo_final.columns.get_loc(coluna) + 1
                letra = ws.cell(row=1,column=col).column_letter
                
                intervalo= f'{letra}2:{letra}{ws.max_row}'
                
                ws.conditional_formatting.add(
                    intervalo, CellIsRule(operator='lessThan', formula=['0'], fill=verde_fill, font=verde_font),
                )
                
                ws.conditional_formatting.add(
                    intervalo, CellIsRule(operator='greaterThan', formula=['0'], fill=vermelho_fill, font=vermelho_font),
                )
                
            # Ajustar largura das colunas
            for coluna in ws.columns:
                maior_tamanho = 0
                letra = coluna[0].column_letter
                
                for celula in coluna:
                    if celula.value is not None:
                        tamanho = len(str(celula.value))
                        maior_tamanho= max(maior_tamanho, tamanho)
                
                ws.column_dimensions[letra].width = maior_tamanho + 2
                
    output.seek(0)
    
    return output